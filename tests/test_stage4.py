"""Stage 4: writing tools and routines (no Windows, AI simulated)."""

import json
from datetime import datetime

import pytest

from agent import config
from agent import memory as memory_module
from agent.agent import resolve_action
from agent.tooling import load_all_tools
from agent.tools import routines, writing
from tests.test_offline import FakeBrain, FakeRouter, FakeUI, call, make_agent


def routine_call(tool_name, **args):
    return {"id": f"id-{tool_name}", "name": tool_name, "arguments": json.dumps(args, ensure_ascii=False)}

TOOLS = load_all_tools()


@pytest.fixture(autouse=True)
def data(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "MEMORY_FILE", tmp_path / "memory.json")
    monkeypatch.setattr(config, "UNDO_FILE", tmp_path / "undo.json")
    monkeypatch.setattr(memory_module, "_instance", None)
    writing._captured.update(text="", hwnd=None, window="", time=0.0, pending=False)


# ---------------------------------------------------------------- Word documents

def test_markdown_becomes_a_formatted_word_document(tmp_path):
    import docx
    from docx.oxml.ns import qn
    md = """# 회의록
참석자: **김철수**, 이영희

## 결정 사항
- 예산 *확정*
- 일정 조정
1. 첫째
2. 둘째

| 항목 | 담당 |
|---|---|
| 디자인 | 김철수 |
"""
    result = TOOLS["create_word_document"].call({"path": str(tmp_path / "minutes"), "content": md,
                                                 "title": "주간 회의", "open_after": False})
    assert result["saved"].endswith("minutes.docx")
    document = docx.Document(result["saved"])
    styles = [(p.style.name, p.text) for p in document.paragraphs if p.text]
    assert ("Title", "주간 회의") in styles and ("Heading 1", "회의록") in styles
    assert ("List Bullet", "예산 확정") in styles and ("List Number", "둘째") in styles
    bold = [r.text for p in document.paragraphs for r in p.runs if r.bold]
    assert "김철수" in bold
    table = document.tables[0]
    assert table.cell(1, 0).text == "디자인" and table.cell(0, 1).paragraphs[0].runs[0].bold
    fonts = document.styles["Normal"].element.rPr.find(qn("w:rFonts"))
    assert fonts.get(qn("w:eastAsia")) == "Malgun Gothic"


def test_word_document_overwrite_asks_first(tmp_path):
    (tmp_path / "a.docx").write_text("x")
    assert TOOLS["create_word_document"].risk_reason({"path": str(tmp_path / "a.docx"), "content": ""})
    assert TOOLS["create_word_document"].risk_reason({"path": str(tmp_path / "b.docx"), "content": ""}) is None


# ---------------------------------------------------------------- email, style, selection

def test_email_draft_opens_prefilled_gmail_and_never_sends(monkeypatch):
    opened = []
    monkeypatch.setattr(writing.webbrowser, "open", opened.append)
    result = TOOLS["open_email_draft"].call({"to": "a@b.com", "subject": "회의 일정", "body": "안녕하세요\n내일 뵙겠습니다."})
    assert opened[0].startswith("https://mail.google.com/mail/?view=cm")
    assert "su=%ED%9A%8C%EC%9D%98" in opened[0] and "%0A" in opened[0]
    assert "sends" in result["note"]


def test_style_guide_is_used_in_the_prompt(tmp_path):
    TOOLS["save_style_guide"].call({"guide": "존댓말(해요체), 짧은 문장, 이모지 없음."})
    agent = make_agent(FakeRouter(FakeBrain([])), FakeUI(), tmp_path)
    assert "짧은 문장" in agent._system_prompt()


def test_hotkey_selection_is_used_once(monkeypatch):
    copies = []
    monkeypatch.setattr(writing, "_copy_selection_from", lambda hwnd=None: copies.append(hwnd) or "새로 선택한 글")
    monkeypatch.setattr("agent.tools.screen._target_window", lambda: {"id": 7, "title": "메모장"})
    writing._captured.update(text="단축키로 잡은 글", hwnd=5, window="Word", time=__import__("time").time(), pending=True)
    assert TOOLS["get_selected_text"].call({})["text"] == "단축키로 잡은 글"
    assert TOOLS["get_selected_text"].call({})["text"] == "새로 선택한 글"
    assert copies == [7]


def test_no_selection_is_reported(monkeypatch):
    monkeypatch.setattr(writing, "_copy_selection_from", lambda hwnd=None: "")
    monkeypatch.setattr("agent.tools.screen._target_window", lambda: {"id": 7, "title": "메모장"})
    assert "select" in TOOLS["get_selected_text"].call({})["error"]


# ---------------------------------------------------------------- routines

def test_schedules_are_parsed_in_korean_english_german():
    assert routines.parse_schedule("9:00", "평일") == {"time": "09:00", "days": [0, 1, 2, 3, 4]}
    assert routines.parse_schedule("18:30", "월,수,금")["days"] == [0, 2, 4]
    assert routines.parse_schedule("7시", "주말") == {"time": "07:00", "days": [5, 6]}
    assert routines.parse_schedule("10:00", "Mo,Mi,Fr")["days"] == [0, 2, 4]
    assert routines.parse_schedule("", "daily") is None
    with pytest.raises(ValueError):
        routines.parse_schedule("25:00", "daily")


def test_steps_routine_and_schedule():
    TOOLS["save_routine"].call({"name": "아침 루틴", "steps": ["날씨 알려주기", "뉴스 요약"], "time": "09:00", "days": "평일"})
    result = TOOLS["run_routine"].call({"name": "아침"})
    assert result["steps"] == ["날씨 알려주기", "뉴스 요약"] and "instruction" in result
    monday_9_05 = datetime(2026, 10, 19, 9, 5)
    due = routines.due_routines(monday_9_05)
    assert due == [{"name": "아침 루틴", "slot": "2026-10-19T09:00"}]  # running it by hand didn't cancel it
    assert routines.due_routines(datetime(2026, 10, 19, 9, 30)) == []    # more than 15 minutes late
    assert routines.due_routines(datetime(2026, 10, 18, 9, 5)) == []     # Sunday
    routines.mark_scheduled("아침 루틴", due[0]["slot"])
    assert routines.due_routines(monday_9_05) == []                     # this slot already ran


def test_schedule_just_before_midnight_still_runs_after_midnight():
    TOOLS["save_routine"].call({"name": "마감", "steps": ["정리"], "time": "23:55", "days": "월"})
    tuesday_00_05 = datetime(2026, 10, 20, 0, 5)
    assert routines.due_routines(tuesday_00_05) == [{"name": "마감", "slot": "2026-10-19T23:55"}]


def test_day_ranges_and_unknown_words():
    assert routines.parse_days("mon-fri") == routines.parse_days("월~금") == [0, 1, 2, 3, 4]
    assert routines.parse_days("월수금") == [0, 2, 4]
    with pytest.raises(ValueError):
        routines.parse_days("someday")


def test_routine_names_must_match_exactly_to_delete():
    TOOLS["save_routine"].call({"name": "아침 루틴", "steps": ["a"]})
    assert "error" in TOOLS["delete_routine"].call({"name": "주말 아침 루틴"})
    assert "error" in TOOLS["delete_routine"].call({"name": ""})
    assert TOOLS["delete_routine"].call({"name": "아침 루틴"})["deleted"] == "아침 루틴"


def test_record_and_replay_a_routine(tmp_path, monkeypatch):
    volume_calls = []
    monkeypatch.setattr(TOOLS["set_volume"], "func", lambda percent: volume_calls.append(percent) or {"volume": percent})
    brain = FakeBrain([
        ("", [call("get_known_folders"), call("set_volume", percent=30)]), ("볼륨 30으로 했어요.", []),
        ("", [routine_call("record_routine", name="조용히")]), ("저장했어요.", []),
        ("", [routine_call("run_routine", name="조용히")]), ("루틴을 실행했어요.", []),
    ])
    agent = make_agent(FakeRouter(brain), FakeUI(), tmp_path)
    agent.handle("볼륨 30으로 해줘")
    agent.handle("방금 한 거 '조용히'로 저장해줘")
    saved = routines.load_routines()[0]
    assert saved["type"] == "recorded" and saved["goal"] == "볼륨 30으로 해줘"
    assert [a["tool"] for a in saved["actions"]] == ["set_volume"]  # looking around is not recorded
    agent.handle("조용히 루틴 실행")
    assert volume_calls == [30, 30]
    replay_result = json.loads(brain.sent[-1][-1]["content"])
    assert replay_result["status"] == "All steps done."
    assert agent.last_task["command"] == "조용히 루틴 실행"  # the replayed actions count as that command's work


def test_replay_keeps_safety_checks_and_hands_back_the_goal(tmp_path):
    victim = tmp_path / "old.txt"
    victim.write_text("x")
    routines._save([{"name": "정리", "type": "recorded", "goal": "오래된 파일 지워줘", "schedule": None,
                     "actions": [{"tool": "delete_item", "args": {"path": str(victim)}, "label": None}]}])
    brain = FakeBrain([("", [routine_call("run_routine", name="정리")]), ("멈췄어요.", [])])
    ui = FakeUI(approve=False)
    make_agent(FakeRouter(brain), ui, tmp_path).handle("정리 루틴 실행")
    assert victim.exists() and ui.confirms
    result = json.loads(brain.sent[-1][-1]["content"])
    assert result["failed_step"] == 1 and result["declined"] and result["goal"] == "오래된 파일 지워줘"
    assert "Do not continue" in result["instruction"]


def test_failed_replay_step_hands_the_goal_to_the_ai(tmp_path):
    routines._save([{"name": "보고서", "type": "recorded", "goal": "보고서 열어줘", "schedule": None,
                     "actions": [{"tool": "read_file", "args": {"path": str(tmp_path / "gone.txt")}, "label": None}]}])
    brain = FakeBrain([("", [routine_call("run_routine", name="보고서")]), ("다른 방법으로 할게요.", [])])
    make_agent(FakeRouter(brain), FakeUI(), tmp_path).handle("보고서 루틴")
    result = json.loads(brain.sent[-1][-1]["content"])
    assert result["failed_step"] == 1 and "Finish the original goal" in result["instruction"]


def test_recorded_clicks_are_found_again_by_name(monkeypatch):
    monkeypatch.setattr("agent.tools.screen._target_window", lambda: {"program": "notepad.exe"})
    assert resolve_action({"tool": "click_element", "args": {"number": 7}, "label": "저장"}) == \
        [("find_and_click", {"name": "저장"})]
    with pytest.raises(LookupError):  # a different program is in front now
        resolve_action({"tool": "click_element", "args": {"number": 7}, "label": "저장", "program": "winword.exe"})
    with pytest.raises(LookupError):  # unnamed elements can't be found again
        resolve_action({"tool": "click_element", "args": {"number": 7}, "label": None})
    steps = resolve_action({"tool": "type_into_element", "args": {"number": 3, "text": "안녕"}, "label": "제목"})
    assert [s[0] for s in steps] == ["find_and_click", "press_keys", "type_text"]
    from agent.tools import web
    monkeypatch.setattr(web, "find_ref", lambda label, occurrence=1: {("뉴스", 1): 12, ("뉴스", 2): 15}.get((label, occurrence)))
    monkeypatch.setattr(web, "current_host", lambda: "news.example.com")
    assert resolve_action({"tool": "web_click", "args": {"number": 4}, "label": "뉴스"}) == [("web_click", {"number": 12})]
    assert resolve_action({"tool": "web_click", "args": {"number": 4}, "label": "뉴스", "occurrence": 2}) == \
        [("web_click", {"number": 15})]
    with pytest.raises(LookupError):
        resolve_action({"tool": "web_click", "args": {"number": 4}, "label": "없는 버튼"})
    with pytest.raises(LookupError):  # recorded on another site
        resolve_action({"tool": "web_click", "args": {"number": 4}, "label": "뉴스", "host": "shop.example.com"})
